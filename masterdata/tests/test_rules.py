"""Бизнес-правила модуля (МД-1…35, §6) сверх критериев приёмки."""
import io

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile

from masterdata import fieldtypes
from masterdata import models as m
from masterdata.validators import iin_bin_checksum_ok, normalize_phone

from .base import MDTestCase


class ValidatorTests(MDTestCase):
    def test_iin_checksum_and_phone(self):
        self.assertTrue(iin_bin_checksum_ok('900101300017'))
        self.assertFalse(iin_bin_checksum_ok('900101300018'))
        self.assertFalse(iin_bin_checksum_ok('12345'))
        self.assertEqual(normalize_phone('8 (701) 111-22-33'), '+77011112233')
        self.assertEqual(normalize_phone('7011112233'), '+77011112233')
        with self.assertRaises(Exception):
            normalize_phone('123')

    def test_fieldtype_coercion(self):
        class F:
            type, multiple, options, ref_directory, validation = 'money', False, [], '', {}
        self.assertEqual(fieldtypes.coerce_json_value(F, '1 500,5 ₸'), 1500.5)
        F.type = 'bool'
        self.assertTrue(fieldtypes.coerce_json_value(F, 'Да'))
        F.type = 'date'
        self.assertEqual(fieldtypes.coerce_json_value(F, '05.03.2026'), '2026-03-05')
        F.type, F.validation = 'string', {'mask': r'^[A-Z]{2}$', 'max_len': 2}
        with self.assertRaises(Exception):
            fieldtypes.coerce_json_value(F, 'abc')
        F.type, F.validation = 'integer', {'min': 1, 'max': 5}
        with self.assertRaises(Exception):
            fieldtypes.coerce_json_value(F, '7')
        F.type, F.multiple, F.options = 'choice', True, [{'value': 'a', 'label': 'Атлас'}, {'value': 'f', 'label': 'Фатин'}]
        self.assertEqual(fieldtypes.coerce_json_value(F, 'атлас; ФАТИН'), ['a', 'f'])


class StructureTests(MDTestCase):
    def test_system_dictionary_rules(self):
        # МД-4: системный нельзя удалить, но можно отключить из меню
        r = self.client.delete('/api/md/directories/models/')
        self.assertEqual(r.status_code, 400)
        self.ok(self.client.patch('/api/md/directories/models/', {'menu_enabled': False}, format='json'))
        menu = self.ok(self.client.get('/api/md/org/'))['menu']
        self.assertFalse([d for d in menu if d['code'] == 'models'][0]['menu_enabled'])
        # МД-10: системному справочнику можно добавить поле (хранится в extra)
        f = self.ok(self.client.post('/api/md/directories/models/fields/', {
            'code': 'silhouette', 'name': 'Силуэт', 'type': 'choice', 'options': ['А-силуэт', 'Пышное'],
            'filterable': True}, format='json'), 201)
        model = self.create_dress(silhouette='Пышное')
        self.assertEqual(model['silhouette'], 'Пышное')
        self.assertEqual(m.ProductModel.objects.get(pk=model['id']).extra['silhouette'], 'Пышное')
        self.assertEqual(self.client.get('/api/md/d/models/', {'silhouette': 'Пышное'}).data['total'], 1)
        # системное поле: подпись менять можно, тип — нельзя, удалить — нельзя
        sys_f = m.DirectoryField.objects.get(directory__organization=self.org, directory__code='models', code='name')
        self.ok(self.client.patch(f'/api/md/directories/models/fields/{sys_f.pk}/', {'name': 'Название модели'},
                                  format='json'))
        self.assertEqual(self.client.patch(f'/api/md/directories/models/fields/{sys_f.pk}/', {'type': 'text'},
                                           format='json').status_code, 400)
        self.assertEqual(self.client.delete(f'/api/md/directories/models/fields/{sys_f.pk}/').status_code, 400)
        # МД-11: удаление поля с данными — только с подтверждением, значения архивируются
        r = self.client.delete(f'/api/md/directories/models/fields/{f["id"]}/')
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.data['requires_confirm'])
        self.ok(self.client.delete(f'/api/md/directories/models/fields/{f["id"]}/?confirm=1'))
        self.assertTrue(m.DirectoryField.objects.get(pk=f['id']).is_archived)
        self.assertEqual(m.ProductModel.objects.get(pk=model['id']).extra['silhouette'], 'Пышное')
        # МД-12: журнал структуры
        log = self.ok(self.client.get('/api/md/directories/models/log/'))
        self.assertIn('field_archive', [x['action'] for x in log])

    def test_custom_directory_lifecycle(self):
        self.ok(self.client.post('/api/md/directories/', {'code': 'fabrics', 'name': 'Ткани', 'is_hierarchical': True},
                                 format='json'), 201)
        self.assertEqual(self.client.post('/api/md/directories/', {'code': 'fabrics', 'name': 'x'},
                                          format='json').status_code, 400)
        self.assertEqual(self.client.post('/api/md/directories/', {'code': 'Ткани', 'name': 'x'},
                                          format='json').status_code, 400)
        cnt = self.ok(self.client.post('/api/md/directories/fabrics/fields/', {'code': 'width', 'name': 'Ширина',
                                                                               'type': 'integer'}, format='json'), 201)
        # МД-11: безопасная смена типа без данных и с данными
        parent = self.ok(self.client.post('/api/md/d/fabrics/', {'name': 'Натуральные', 'width': 150}, format='json'), 201)
        child = self.ok(self.client.post('/api/md/d/fabrics/', {'name': 'Шёлк', 'parent': parent['id']}, format='json'), 201)
        self.assertEqual(child['parent__display'], 'Натуральные')
        self.ok(self.client.patch(f'/api/md/directories/fabrics/fields/{cnt["id"]}/', {'type': 'decimal'}, format='json'))
        self.assertEqual(self.client.patch(f'/api/md/directories/fabrics/fields/{cnt["id"]}/', {'type': 'date'},
                                           format='json').status_code, 400)
        # МД-3: копия структуры без данных
        cp = self.ok(self.client.post('/api/md/directories/fabrics/copy/', {'code': 'furnitura', 'name': 'Фурнитура'},
                                      format='json'), 201)
        self.assertEqual(cp['code'], 'furnitura')
        self.assertEqual(self.client.get('/api/md/d/furnitura/').data['total'], 0)
        self.assertEqual(m.DirectoryField.objects.filter(directory__code='furnitura').count(), 2)
        # МД-5: справочник с записями удаляется только в архив; пустой — удаляется
        r = self.ok(self.client.delete('/api/md/directories/fabrics/'))
        self.assertEqual(r['result'], 'archived')
        r = self.ok(self.client.delete('/api/md/directories/furnitura/'))
        self.assertEqual(r['result'], 'deleted')
        # ссылка из системного справочника на пользовательский
        self.ok(self.client.post('/api/md/directories/fabrics/restore/'))
        self.ok(self.client.post('/api/md/directories/models/fields/', {
            'code': 'fabric', 'name': 'Ткань', 'type': 'ref', 'ref_directory': 'fabrics'}, format='json'), 201)
        model = self.create_dress(fabric='Шёлк')
        self.assertEqual(model['fabric__display'], 'Шёлк')
        # на «Шёлк» есть ссылка — удаление записи переводит её в архив (МД-21)
        r = self.ok(self.client.delete(f'/api/md/d/fabrics/{child["id"]}/'))
        self.assertEqual(r['result'], 'archived')

    def test_only_manager_changes_structure(self):
        seller = self.api(self.make_user('s@test.kz', 'seller'))
        self.assertEqual(seller.post('/api/md/directories/', {'code': 'x1', 'name': 'X'}, format='json').status_code, 403)
        self.assertEqual(seller.patch('/api/md/org/', {'sections': {'uzatu': False}}, format='json').status_code, 403)
        self.assertEqual(seller.get('/api/md/role-policies/').status_code, 403)


class RecordRuleTests(MDTestCase):
    def test_archive_restore_and_delete(self):
        ev = self.ok(self.client.post('/api/md/d/events/', {'name': 'Наурыз'}, format='json'), 201)
        # без ссылок — удаляется физически
        self.assertEqual(self.ok(self.client.delete(f'/api/md/d/events/{ev["id"]}/'))['result'], 'deleted')
        ev = self.ok(self.client.post('/api/md/d/events/', {'name': 'Наурыз'}, format='json'), 201)
        self.create_dress(events=[ev['id']])
        self.assertEqual(self.ok(self.client.delete(f'/api/md/d/events/{ev["id"]}/'))['result'], 'archived')
        # архивные скрыты из списка и выбора, но видны с «Показать архив»
        self.assertNotIn(ev['id'], [o['id'] for o in self.ok(self.client.get('/api/md/d/events/options/'))])
        self.assertNotIn(ev['id'], [x['id'] for x in self.client.get('/api/md/d/events/').data['data']])
        self.assertIn(ev['id'], [x['id'] for x in self.client.get('/api/md/d/events/', {'archived': '1'}).data['data']])
        self.ok(self.client.post(f'/api/md/d/events/{ev["id"]}/restore/'))
        self.assertIn(ev['id'], [o['id'] for o in self.ok(self.client.get('/api/md/d/events/options/'))])

    def test_category_tree_rules(self):
        root = self.cat('Ұзату')
        lvl2 = self.ok(self.client.post('/api/md/d/categories/', {'name': 'Головные уборы', 'parent': self.cat('Тақия')},
                                        format='json'), 201)
        r = self.client.post('/api/md/d/categories/', {'name': '4-й уровень', 'parent': lvl2['id']}, format='json')
        self.assertEqual(r.status_code, 400)  # глубина до 3 уровней
        r = self.client.patch(f'/api/md/d/categories/{root}/', {'parent': lvl2['id']}, format='json')
        self.assertEqual(r.status_code, 400)  # цикл
        # нельзя добавить подкатегорию к категории с моделями
        self.create_dress()
        r = self.client.post('/api/md/d/categories/', {'name': 'Мини', 'parent': self.cat('Свадебные платья')},
                             format='json')
        self.assertEqual(r.status_code, 400)
        # категория с префиксом и своей маской артикула {PREFIX}-{COUNTRY}-{SEQ:4}
        cat = self.ok(self.client.post('/api/md/d/categories/', {
            'name': 'Вечерние', 'article_prefix': 'вч', 'operations': 'rent', 'model_mask': '{PREFIX}-{COUNTRY}-{SEQ:4}'},
            format='json'), 201)
        self.assertEqual(cat['article_prefix'], 'ВЧ')
        model = self.ok(self.client.post('/api/md/d/models/', {
            'name': 'Вечернее', 'category': cat['id'], 'events': [self.event('Прочие')], 'rent_price': 50000,
            'country': self.get_id(m.Country, iso_code='TR')}, format='json'), 201)
        self.assertEqual(model['article'], 'ВЧ-TR-0001')

    def test_item_rules(self):
        model = self.create_dress(acquisition='purchase')
        r = self.client.post('/api/md/d/items/', {'model': model['id'], 'salon': self.salon()}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('size', r.data['fields'])
        self.assertIn('supplier', r.data['fields'])
        item = self.create_item(model['id'], '44', barcode='4600000000001')
        r = self.client.post('/api/md/d/items/', {'model': model['id'], 'size': self.size('42'), 'salon': self.salon(),
                                                  'supplier': self.supplier(), 'barcode': '4600000000001'}, format='json')
        self.assertIn('barcode', r.data['fields'])
        # бирка неизменна, число аренд считает система
        r = self.client.patch(f'/api/md/d/items/{item["id"]}/', {'tag': '999999', 'rent_count': 5}, format='json')
        self.assertEqual(r.status_code, 400)
        it = self.ok(self.client.get(f'/api/md/d/items/{item["id"]}/'))
        self.assertEqual(it['effective_prices']['rent_price'], 150000)
        self.assertEqual(it['resource_left'], 10)
        # индивидуальная цена переопределяет цену модели
        it = self.ok(self.client.patch(f'/api/md/d/items/{item["id"]}/', {'rent_price_override': 120000}, format='json'))
        self.assertEqual(it['effective_prices']['rent_price'], 120000)
        # фильтры товара: категория (с подкатегориями), свободен, стадия
        self.assertEqual(self.client.get('/api/md/d/items/', {'category': self.cat('Свадебные платья')}).data['total'], 1)
        self.assertEqual(self.client.get('/api/md/d/items/', {'available': 'true'}).data['total'], 1)
        self.assertEqual(self.client.get('/api/md/d/models/', {'size': self.size('44')}).data['total'], 1)
        self.assertEqual(self.client.get('/api/md/d/models/', {'size': self.size('48')}).data['total'], 0)

    def test_counterparty_rules(self):
        r = self.client.post('/api/md/d/counterparties/', {'name': 'ТОО', 'kind': 'too', 'roles': []}, format='json')
        self.assertIn('roles', r.data['fields'])
        r = self.client.post('/api/md/d/counterparties/', {'name': 'ТОО', 'kind': 'too', 'roles': ['dry_cleaning'],
                                                           'bin_iin': '123456789012'}, format='json')
        self.assertIn('bin_iin', r.data['fields'])
        cp = self.ok(self.client.post('/api/md/d/counterparties/', {
            'name': 'Химчистка «Ақ»', 'kind': 'ip', 'roles': ['dry_cleaning', 'restorer'], 'email': 'ak@mail.kz'},
            format='json'), 201)
        self.assertEqual(cp['roles__display'], ['Химчистка', 'Реставратор'])
        self.assertEqual(self.client.get('/api/md/d/counterparties/', {'roles': 'restorer'}).data['total'], 1)


class BulkExportTests(MDTestCase):
    def test_bulk_export_template_prefs(self):
        a = self.create_dress(name='A')
        b = self.create_dress(name='B')
        r = self.ok(self.client.post('/api/md/d/models/bulk/', {
            'ids': [a['id'], b['id']], 'action': 'update', 'field': 'show_in_catalog', 'value': True}, format='json'))
        self.assertEqual(r['done'], 2)
        self.assertEqual(self.client.get('/api/md/d/models/', {'show_in_catalog': 'true'}).data['total'], 2)
        r = self.ok(self.client.post('/api/md/d/models/bulk/', {
            'ids': [a['id']], 'action': 'update', 'field': 'min_rent_price', 'value': 999999}, format='json'))
        self.assertEqual(len(r['errors']), 1)  # правила работают и в массовых действиях
        self.ok(self.client.post('/api/md/d/models/bulk/', {'ids': [b['id']], 'action': 'archive'}, format='json'))
        self.assertEqual(self.client.get('/api/md/d/models/').data['total'], 1)
        # экспорт XLSX/CSV текущего списка с фильтрами
        x = self.client.get('/api/md/d/models/export/', {'format': 'xlsx', 'archived': '1'})
        wb = openpyxl.load_workbook(io.BytesIO(x.content))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(len(rows), 3)
        self.assertIn('Цена аренды, ₸', rows[0])
        c = self.client.get('/api/md/d/models/export/', {'format': 'csv', 'q': 'A'})
        text = c.content.decode('utf-8-sig')
        self.assertEqual(len(text.strip().splitlines()), 2)
        # шаблон импорта: обязательные отмечены, выпадающие списки, инструкция
        t = self.client.get('/api/md/d/models/template/')
        wb = openpyxl.load_workbook(io.BytesIO(t.content))
        self.assertEqual(wb.sheetnames, ['Данные', 'Списки', 'Инструкция'])
        hdr = [c.value for c in wb['Данные'][1]]
        self.assertIn('Наименование *', hdr)
        self.assertTrue(len(wb['Данные'].data_validations.dataValidation) >= 3)
        # настройки колонок пользователя
        self.ok(self.client.put('/api/md/prefs/columns:models/', {'value': {'columns': ['article', 'name']}},
                                format='json'))
        self.assertEqual(self.ok(self.client.get('/api/md/d/models/meta/'))['columns_pref']['columns'],
                         ['article', 'name'])

    def test_items_labels_bulk(self):
        model = self.create_dress()
        i1 = self.create_item(model['id'], '42')
        i2 = self.create_item(model['id'], '44')
        r = self.client.post('/api/md/d/items/bulk/', {'ids': [i1['id'], i2['id']], 'action': 'labels'}, format='json')
        self.assertEqual(r['Content-Type'], 'application/pdf')


class RolesAndImportTests(MDTestCase):
    def test_accountant_rights(self):
        acc = self.api(self.make_user('acc@test.kz', 'accountant'))
        model = self.create_dress()
        item = self.create_item(model['id'], '44', cost=250000)
        # видит закупку, не меняет товар
        it = self.ok(acc.get(f'/api/md/d/items/{item["id"]}/'))
        self.assertEqual(it['cost'], 250000)
        self.assertEqual(acc.patch(f'/api/md/d/models/{model["id"]}/', {'name': 'x'}, format='json').status_code, 403)
        # счета и контрагенты — да, импорт только их
        self.ok(acc.post('/api/md/d/accounts/', {'name': 'Аренда склада', 'pnl': 'expense', 'cashflow': 'outflow'},
                         format='json'), 201)
        up = acc.post('/api/md/import/upload/', {'directory': 'models', 'file': SimpleUploadedFile('a.csv', b'a;b\n1;2\n')},
                      format='multipart')
        self.assertEqual(up.status_code, 403)
        csv_data = 'Наименование;Тип;Роли;ИИН / БИН\nТОО «Ткани Востока»;ТОО;Ткани и фурнитура;\nИП Швея;ИП;"Ателье / швея; Вышивка";\n'
        up = self.ok(acc.post('/api/md/import/upload/', {
            'directory': 'counterparties', 'file': SimpleUploadedFile('cp.csv', csv_data.encode())}, format='multipart'), 201)
        job = up['job']['id']
        prev = self.ok(acc.post(f'/api/md/import/{job}/preview/', {'mapping': up['mapping']}, format='json'))
        self.assertEqual(prev['counts']['create'], 2, prev['errors'])
        self.ok(acc.post(f'/api/md/import/{job}/run/', {}, format='json'))
        self.assertEqual(m.Counterparty.objects.filter(organization=self.org).count(), 3)  # +поставщик из create_item
        # администратору счета — только просмотр
        admin = self.api(self.make_user('adm@test.kz', 'admin'))
        self.assertEqual(admin.post('/api/md/d/accounts/', {'name': 'x', 'pnl': 'none', 'cashflow': 'none'},
                                    format='json').status_code, 403)
        # продавцу счета не видны вообще
        seller = self.api(self.make_user('s@test.kz', 'seller'))
        self.assertEqual(seller.get('/api/md/d/accounts/').status_code, 403)

    def test_role_matrix_edit(self):
        seller_u = self.make_user('s@test.kz', 'seller')
        seller = self.api(seller_u)
        self.assertEqual(seller.get('/api/md/d/models/export/').status_code, 403)
        self.ok(self.client.put('/api/md/role-policies/', {'rows': [
            {'role': 'seller', 'directory_code': 'models', 'can_export': True, 'hidden_fields': ['sale_price']}]},
            format='json'))
        self.assertEqual(seller.get('/api/md/d/models/export/').status_code, 200)
        model = self.create_dress()
        self.assertNotIn('sale_price', self.ok(seller.get(f'/api/md/d/models/{model["id"]}/')))

    def test_import_update_mode_and_profiles(self):
        c = self.ok(self.client.post('/api/md/d/clients/', {'full_name': 'Анна', 'phone': '+77011112233'},
                                     format='json'), 201)
        data = 'ФИО;Телефон;Город\nАнна Иванова;8 701 111 22 33;Астана\nНовая;+7 702 000 00 00;Шымкент\n'
        up = self.ok(self.client.post('/api/md/import/upload/', {
            'directory': 'clients', 'file': SimpleUploadedFile('c.csv', data.encode())}, format='multipart'), 201)
        job = up['job']['id']
        prev = self.ok(self.client.post(f'/api/md/import/{job}/preview/', {
            'mapping': up['mapping'], 'mode': 'update', 'key_field': 'phone', 'profile_name': 'Клиенты'}, format='json'))
        self.assertEqual(prev['counts'], {'create': 0, 'update': 1, 'skip': 1, 'error': 0})
        self.ok(self.client.post(f'/api/md/import/{job}/run/', {}, format='json'))
        got = self.ok(self.client.get(f'/api/md/d/clients/{c["id"]}/'))
        self.assertEqual(got['city'], 'Астана')
        self.assertEqual(got['full_name'], 'Анна Иванова')
        hist = self.ok(self.client.get(f'/api/md/d/clients/{c["id"]}/history/'))
        self.assertEqual(hist[0]['source'], 'import')
        # профиль сопоставления сохранён и предлагается при следующей загрузке
        up2 = self.ok(self.client.post('/api/md/import/upload/', {
            'directory': 'clients', 'file': SimpleUploadedFile('c.csv', data.encode())}, format='multipart'), 201)
        self.assertEqual(up2['profiles'][0]['name'], 'Клиенты')
        # создание недостающих значений справочников (выключено по умолчанию)
        data = 'Наименование;Категория;Мероприятия;Цена аренды, ₸;Цена продажи, ₸\nПлатье X;Свадебные платья;Тойбастар;100000;500000\n'
        up = self.ok(self.client.post('/api/md/import/upload/', {
            'directory': 'models', 'file': SimpleUploadedFile('m.csv', data.encode())}, format='multipart'), 201)
        prev = self.ok(self.client.post(f'/api/md/import/{up["job"]["id"]}/preview/', {'mapping': up['mapping']},
                                        format='json'))
        self.assertEqual(prev['counts']['error'], 1)
        prev = self.ok(self.client.post(f'/api/md/import/{up["job"]["id"]}/preview/', {
            'mapping': up['mapping'], 'create_missing_refs': True}, format='json'))
        self.assertEqual(prev['counts']['create'], 1)
        self.assertFalse(m.Event.objects.filter(organization=self.org, name='Тойбастар').exists())  # предпросмотр откатан
        self.ok(self.client.post(f'/api/md/import/{up["job"]["id"]}/run/', {}, format='json'))
        self.assertTrue(m.Event.objects.filter(organization=self.org, name='Тойбастар').exists())

    def test_only_valid_rows(self):
        data = 'ФИО;Телефон\nАнна;+77011112233\nБез телефона;\nОльга;12\n'
        up = self.ok(self.client.post('/api/md/import/upload/', {
            'directory': 'clients', 'file': SimpleUploadedFile('c.csv', data.encode())}, format='multipart'), 201)
        job = up['job']['id']
        prev = self.ok(self.client.post(f'/api/md/import/{job}/preview/', {'mapping': up['mapping']}, format='json'))
        self.assertEqual(prev['counts']['error'], 2)
        run = self.ok(self.client.post(f'/api/md/import/{job}/run/', {'only_valid': True}, format='json'))
        self.assertEqual(run['status'], 'done')
        self.assertEqual(m.Client.objects.filter(organization=self.org).count(), 1)
        # журнал и файл ошибок
        self.assertEqual(self.client.get(f'/api/md/import/{job}/errors-file/').status_code, 200)


class PhotoAndOrgTests(MDTestCase):
    def test_photo_upload_rules(self):
        from .test_acceptance import jpeg_bytes
        model = self.create_dress()
        r = self.client.post(f'/api/md/d/models/{model["id"]}/photos/', {
            'files': [SimpleUploadedFile('a.jpg', jpeg_bytes()), SimpleUploadedFile('b.png', jpeg_bytes())]},
            format='multipart')
        photos = self.ok(r, 201)
        self.assertEqual(len(photos), 2)
        r = self.client.post(f'/api/md/d/models/{model["id"]}/photos/', {
            'files': [SimpleUploadedFile('doc.pdf', b'%PDF')]}, format='multipart')
        self.assertEqual(r.status_code, 400)
        second = photos[1]['id']
        self.ok(self.client.patch(f'/api/md/photos/{second}/', {'is_main': True,
                                                               'photo_type': self.get_id(m.PhotoType, name='Спина')},
                                  format='json'))
        got = self.ok(self.client.get(f'/api/md/d/models/{model["id"]}/'))
        self.assertEqual(got['main_photo']['id'], second)
        self.ok(self.client.post('/api/md/photos/reorder/', {'ids': [second, photos[0]['id']]}, format='json'))
        self.ok(self.client.delete(f'/api/md/photos/{second}/'))
        self.assertTrue(m.Photo.objects.get(pk=photos[0]['id']).is_main)

    def test_organization_self_registration(self):
        from accounts.models import User
        u = User.objects.create_user(email='new@salon.kz', name='Новый', phone='+77010000009')
        u.status = 'active'
        u.save()
        c = self.api(u)
        self.assertEqual(c.get('/api/md/org/').status_code, 403)
        self.ok(c.post('/api/md/organizations/', {'name': 'Салон «Арман»'}, format='json'), 201)
        me = self.ok(c.get('/api/md/org/'))
        self.assertEqual(me['role'], 'owner')
        self.assertTrue(any(d['code'] == 'models' for d in me['menu']))
        # участники
        self.ok(self.client.post('/api/md/members/', {'email': 'new@salon.kz', 'role': 'seller'}, format='json'))
        self.assertEqual(len(self.ok(self.client.get('/api/md/members/'))), 2)

    def test_openapi_and_pages(self):
        r = self.client.get('/api/md/schema/')
        self.assertEqual(r.status_code, 200)
        self.assertIn('/api/md/d/{code}/', r.json()['paths'])
        self.assertEqual(self.client.get('/md').status_code, 200)
        self.assertEqual(self.client.get('/md-assets/md-app.js').status_code, 200)
