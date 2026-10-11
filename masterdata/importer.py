"""
Импорт данных (МД-25…35).

Шаги мастера: файл → сопоставление колонок → предпросмотр и ошибки → результат.
Каждая строка проходит через тот же движок (backends), что и ручной ввод,
поэтому все правила (МД-20) и журнал изменений действуют и при импорте.

Предпросмотр выполняет загрузку целиком внутри транзакции и откатывает её:
так проверяются и правила, и уникальность, и зависимости между строками
(например, модель из 1-й строки используется во 2-й).
"""
import csv
import io
import logging
import re

from django.conf import settings
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from django.utils import timezone
from rest_framework import exceptions, serializers

from . import fieldtypes
from . import models as m
from .articles import bump_counter
from .backends import get_backend
from .refs import resolve_ref

log = logging.getLogger('masterdata')


class Column:
    def __init__(self, key, label, fdef=None, required=False, hint=''):
        self.key, self.label, self.fdef, self.required, self.hint = key, label, fdef, required, hint

    def as_dict(self):
        f = self.fdef
        return {'key': self.key, 'label': self.label, 'required': self.required,
                'type': f.type if f else 'string', 'multiple': bool(f and f.multiple),
                'ref_directory': f.ref_directory if f else '', 'hint': self.hint or (f.hint if f else '')}


class _FakeField:
    def __init__(self, code, name, type='string', **kw):
        self.code, self.name, self.type = code, name, type
        self.multiple = kw.get('multiple', False)
        self.options = kw.get('options', [])
        self.ref_directory = kw.get('ref_directory', '')
        self.validation = kw.get('validation', {})
        self.required = kw.get('required', False)
        self.hint = kw.get('hint', '')
        self.is_system = True


# Ключи сопоставления (МД-28): код, артикул, бирка, ИИН/БИН, телефон
KEY_FIELDS = {
    'models': ['article', 'code'], 'items': ['tag', 'barcode'], 'kits': ['code'],
    'clients': ['iin', 'phone', 'code'], 'counterparties': ['bin_iin', 'code', 'name'],
    'countries': ['iso_code'], 'sizes': ['code', 'value'],
}
KEY_LABELS = {'code': 'Код', 'article': 'Артикул', 'tag': 'Бирка', 'barcode': 'Штрихкод', 'iin': 'ИИН',
              'bin_iin': 'ИИН/БИН', 'phone': 'Телефон', 'name': 'Наименование', 'iso_code': 'Код ISO',
              'value': 'Значение'}

ITEM_MODEL_FIELDS = ['article', 'name', 'category', 'events', 'styles', 'operations', 'acquisition',
                     'manufacturer', 'color', 'rent_price', 'min_rent_price', 'sale_price', 'resource', 'description']


def key_fields(backend):
    keys = KEY_FIELDS.get(backend.code)
    if keys:
        return keys
    return ['code', backend.directory.display_field or 'name']


def columns_for(backend):
    cols = []
    hidden = backend.policy.hidden if backend.policy else set()
    for f in backend.all_fields():
        if f.code in hidden or f.type in ('computed', 'photo', 'file'):
            continue
        if f.readonly and not (backend.code == 'items' and f.code in ('tag', 'rent_count')):
            continue
        if backend.code == 'items' and f.code == 'model':
            continue
        req = f.required and not (backend.code == 'models' and f.code == 'article')
        if backend.code == 'items' and f.code in ('stage', 'received_at'):
            req = False  # по умолчанию: «В салоне», сегодня
        cols.append(Column(f.code, f.name, f, req))
    if backend.code == 'items':
        mb = get_backend(backend.org, 'models', backend.user)
        for code in ITEM_MODEL_FIELDS:
            f = mb.field(code)
            if f is None:
                continue
            cols.append(Column(f'model.{code}', f'Модель: {f.name}', f, code == 'article',
                               hint='Модель находится по артикулу или создаётся' if code == 'article' else ''))
    if backend.hierarchical:
        cols.append(Column('parent', 'Родитель', _FakeField('parent', 'Родитель', 'ref', ref_directory=backend.code),
                           hint='Код или название родительской записи'))
    if not any(c.key == 'code' for c in cols) and backend.code not in ('models', 'items'):
        cols.append(Column('code', 'Код', _FakeField('code', 'Код')))
    return cols


def _norm(s):
    s = str(s or '').strip().lower().replace('ё', 'е').replace('*', '')
    return re.sub(r'[\s_:,.()₸/\\-]+', ' ', s).strip()


def auto_map(headers, columns):
    """МД-27: автоматически по названию колонки (или коду поля)."""
    by_label = {}
    for c in columns:
        by_label.setdefault(_norm(c.label), c.key)
        by_label.setdefault(_norm(c.key), c.key)
    mapping, used = {}, set()
    for h in headers:
        k = by_label.get(_norm(h))
        if k and k not in used:
            mapping[h] = k
            used.add(k)
    for h in headers:  # частичное совпадение для оставшихся
        if h in mapping or not _norm(h):
            continue
        for c in columns:
            if c.key in used:
                continue
            a, b = _norm(h), _norm(c.label)
            if a and b and (a in b or b in a) and min(len(a), len(b)) >= 4:
                mapping[h] = c.key
                used.add(c.key)
                break
    return mapping


# ------------------------------------------------------------
# Чтение файла (МД-26)
# ------------------------------------------------------------
def read_table(raw: bytes, filename: str):
    name = (filename or '').lower()
    if name.endswith('.xlsx') or raw[:2] == b'PK':
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise exceptions.ValidationError({'file': ['Не удалось прочитать XLSX-файл']})
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif name.endswith('.csv') or name.endswith('.txt'):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise exceptions.ValidationError({'file': ['CSV должен быть в кодировке UTF-8']})
        first = text.split('\n', 1)[0]
        delim = ';' if first.count(';') >= first.count(',') else ','
        rows = [r for r in csv.reader(io.StringIO(text), delimiter=delim)]
    else:
        raise exceptions.ValidationError({'file': ['Принимаются файлы XLSX и CSV']})
    # заголовок — первая непустая строка
    while rows and all(fieldtypes.is_empty(c) for c in rows[0]):
        rows.pop(0)
    if not rows:
        raise exceptions.ValidationError({'file': ['Файл пуст']})
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    data = []
    for i, r in enumerate(rows[1:], start=2):
        if all(fieldtypes.is_empty(c) for c in r):
            continue
        data.append((i, list(r) + [None] * (len(headers) - len(r))))
    if len(data) > settings.MD_IMPORT_MAX_ROWS:
        raise exceptions.ValidationError({'file': [f'Не больше {settings.MD_IMPORT_MAX_ROWS} строк за раз']})
    return headers, data


def _cell(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        return v.strip()
    return v


# ------------------------------------------------------------
# Обработка
# ------------------------------------------------------------
def _flatten_errors(detail):
    out = {}
    if isinstance(detail, dict):
        for k, v in detail.items():
            if isinstance(v, dict):
                v = [str(x) for vv in v.values() for x in (vv if isinstance(vv, list) else [vv])]
            out[k] = [str(x) for x in (v if isinstance(v, list) else [v])]
    elif isinstance(detail, list):
        out['non_field_errors'] = [str(x) for x in detail]
    else:
        out['non_field_errors'] = [str(detail)]
    return out


def _convert(org, user, col, raw, create_missing, ref_cache):
    f = col.fdef
    if col.key == 'code' or f is None:
        return str(raw).strip()
    if f.type == 'iin_bin' and isinstance(raw, (int, float)):
        raw = str(int(raw)).zfill(12)

    def cached_resolve(code, r):
        key = (code, str(r).strip().lower())
        if key not in ref_cache:
            ref_cache[key] = resolve_ref(org, code, r, create_missing, user)
        return ref_cache[key]
    return fieldtypes.coerce_json_value(f, raw, resolve_ref=cached_resolve)


def process(job, apply=False, progress=None):
    """Выполняет проверку (apply=False) или загрузку (apply=True)."""
    from .api.context import Policy
    from accounts.models import Membership
    org, user = job.organization, job.user
    mem = Membership.objects.filter(user=user, organization=org).first()
    role = mem.role if mem else 'seller'
    policy = Policy(role, job.directory_code,
                    m.RolePolicy.objects.filter(organization=org, role=role, directory_code=job.directory_code).first())
    backend = get_backend(org, job.directory_code, user, policy)
    cols = {c.key: c for c in columns_for(backend)}
    raw = job.file.open('rb').read()
    job.file.close()
    headers, rows = read_table(raw, job.original_name)
    col_by_idx = {}
    for i, h in enumerate(headers):
        key = (job.mapping or {}).get(h)
        if key and key in cols:
            col_by_idx[i] = cols[key]
    opts = job.options or {}
    create_missing = bool(opts.get('create_missing_refs'))
    only_valid = bool(opts.get('only_valid'))
    mode = job.mode or 'create'
    key = job.key_field or key_fields(backend)[0]
    header_of = {c.key: headers[i] for i, c in col_by_idx.items()}

    counts = {'create': 0, 'update': 0, 'skip': 0, 'error': 0}
    ref_cache, model_cache = {}, {}
    errors, skipped, error_rows = [], [], []
    total = len(rows)
    model_backend = None
    if backend.code == 'items':
        model_policy = Policy(role, 'models', m.RolePolicy.objects.filter(
            organization=org, role=role, directory_code='models').first())
        model_backend = get_backend(org, 'models', user, model_policy)

    def row_error(rownum, field, msgs, row):
        column = header_of.get(field, header_of.get(f'model.{field}', ''))
        if field == 'non_field_errors':
            column = ''
        for msg in msgs:
            errors.append({'row': rownum, 'column': column, 'field': field, 'message': msg})

    with transaction.atomic():
        for n, (rownum, row) in enumerate(rows):
            payload, model_payload, row_errs = {}, {}, {}
            for i, col in col_by_idx.items():
                v = _cell(row[i] if i < len(row) else None)
                if fieldtypes.is_empty(v):
                    continue
                try:
                    val = _convert(org, user, col, v, create_missing, ref_cache)
                except DjangoValidationError as e:
                    row_errs[col.key] = e.messages
                    continue
                except exceptions.ValidationError as e:
                    row_errs[col.key] = _flatten_errors(e.detail).get('non_field_errors', ['Ошибка'])
                    continue
                if col.key.startswith('model.'):
                    model_payload[col.key[6:]] = val
                else:
                    payload[col.key] = val
            if row_errs:
                counts['error'] += 1
                for f, msgs in row_errs.items():
                    row_error(rownum, f, msgs, row)
                error_rows.append((rownum, row, row_errs))
                continue
            sid = transaction.savepoint()
            try:
                instance = None
                kv = payload.get(key)
                if kv not in (None, ''):
                    instance = _find_existing(backend, key, kv)
                if instance is not None and mode == 'create':
                    counts['skip'] += 1
                    skipped.append({'row': rownum, 'reason': f'Уже существует ({KEY_LABELS.get(key, key)}: {kv})'})
                    transaction.savepoint_commit(sid)
                    continue
                if instance is None and mode == 'update':
                    counts['skip'] += 1
                    skipped.append({'row': rownum, 'reason': 'Запись для обновления не найдена'})
                    transaction.savepoint_commit(sid)
                    continue
                if backend.code == 'items':
                    if instance is None or model_payload:
                        model = _resolve_item_model(model_backend, model_payload, instance, create_missing, model_cache)
                        payload['model'] = str(model.pk)
                    if instance is None and payload.get('tag'):
                        _check_tag(org, payload['tag'])
                obj, _w = backend.save(payload, instance, partial=instance is not None, source='import',
                                       create_missing=create_missing)
                transaction.savepoint_commit(sid)
                counts['update' if instance is not None else 'create'] += 1
            except (serializers.ValidationError, exceptions.ValidationError) as e:
                transaction.savepoint_rollback(sid)
                counts['error'] += 1
                fe = _flatten_errors(e.detail)
                for f, msgs in fe.items():
                    row_error(rownum, f, msgs, row)
                error_rows.append((rownum, row, fe))
            except DjangoValidationError as e:
                transaction.savepoint_rollback(sid)
                counts['error'] += 1
                row_error(rownum, 'non_field_errors', e.messages, row)
                error_rows.append((rownum, row, {'non_field_errors': e.messages}))
            if progress and (n % 25 == 0 or n == total - 1):
                progress(n + 1, total)
        failed = apply and counts['error'] > 0 and not only_valid
        if not apply or failed:
            transaction.set_rollback(True)

    job.headers = headers
    job.total_rows = total
    job.processed_rows = total
    job.counts = counts
    job.errors = errors[:2000]
    job.options = {**opts, 'skipped': skipped[:500]}
    if error_rows:
        job.error_file.save(f'errors_{job.original_name.rsplit(".", 1)[0][:60]}.xlsx',
                            ContentFile(_error_file(headers, error_rows, header_of)), save=False)
    elif job.error_file:
        job.error_file = None
    if not apply:
        job.status = 'previewed'
        job.message = ''
    elif failed:
        job.status = 'failed'
        job.message = (f'Импорт не выполнен: ошибки в {counts["error"]} строках. Исправьте файл или выберите '
                       f'«Загрузить только корректные строки».')
        job.counts = {**counts, 'create': 0, 'update': 0}
        job.finished_at = timezone.now()
    else:
        job.status = 'done'
        job.message = f'Создано: {counts["create"]}, обновлено: {counts["update"]}, ' \
                      f'пропущено: {counts["skip"]}, с ошибками: {counts["error"]}'
        job.finished_at = timezone.now()
    job.save()
    return job


def _find_existing(backend, key, value):
    qs = backend.base_qs()
    if backend.code == 'clients' and key == 'iin':
        from .crypto import blind_hash
        return qs.filter(iin_hash=blind_hash(str(value))).first()
    if backend.code == 'clients' and key == 'phone':
        from .validators import phone_digits
        return qs.filter(phone_digits=phone_digits(value)).first()
    if hasattr(backend, 'spec'):
        try:
            return qs.filter(**{f'{key}__iexact': str(value)}).first()
        except Exception:
            return qs.filter(**{key: value}).first()
    if key == 'code':
        return qs.filter(code__iexact=str(value)).first()
    return qs.filter(**{f'data__{key}__iexact': str(value)}).first()


def _resolve_item_model(model_backend, model_payload, instance, create_missing, cache):
    """МД-33: модель находится по артикулу или создаётся."""
    art = (model_payload.get('article') or '').strip() if model_payload.get('article') else ''
    if not art:
        if instance is not None:
            return instance.model
        raise serializers.ValidationError({'article': ['Укажите артикул модели']})
    key = art.lower()
    if key in cache and cache[key].pk and \
            m.ProductModel.objects.filter(pk=cache[key].pk).exists():  # не откатилась ли строка-создатель
        return cache[key]
    model = m.ProductModel.objects.filter(organization=model_backend.org, article__iexact=art) \
        .select_related('category').first()
    if model is None:
        model, _ = model_backend.save(model_payload, None, source='import', create_missing=create_missing)
    cache[key] = model
    return model


def _check_tag(org, tag):
    tag = str(tag).strip()
    if m.Item.objects.filter(organization=org, tag=tag).exists():
        raise serializers.ValidationError({'tag': [f'Бирка {tag} уже есть']})
    if tag.isdigit():
        bump_counter(org, 'TAG', int(tag))


def _error_file(headers, error_rows, header_of):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ошибки'
    ws.append(['Строка'] + headers + ['Ошибки'])
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill('solid', fgColor='FFE0E0')
    inv = {v: k for k, v in header_of.items()}
    for rownum, row, errs in error_rows:
        msgs = []
        for f, ms in errs.items():
            col = header_of.get(f) or header_of.get(f'model.{f}') or ''
            msgs.extend(f'{col}: {x}' if col else x for x in ms)
        ws.append([rownum] + [v for v in row[:len(headers)]] + ['; '.join(msgs)])
        r = ws.max_row
        for f in errs:
            col = header_of.get(f) or header_of.get(f'model.{f}')
            if col in headers:
                ws.cell(row=r, column=headers.index(col) + 2).fill = red
    del inv
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------
# Шаблон XLSX (МД-25)
# ------------------------------------------------------------
def build_template(backend):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from .refs import ref_queryset
    cols = columns_for(backend)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Данные'
    lists = wb.create_sheet('Списки')
    help_ws = wb.create_sheet('Инструкция')
    req_fill = PatternFill('solid', fgColor='FCE4D6')
    hdr_fill = PatternFill('solid', fgColor='DDEBF7')
    list_col = 0
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c.label + (' *' if c.required else ''))
        cell.font = Font(bold=True)
        cell.fill = req_fill if c.required else hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(c.label) + 4))
        f = c.fdef
        values = None
        if f is not None and not f.multiple:
            if f.type == 'choice':
                values = [o['label'] for o in f.options]
            elif f.type == 'bool':
                values = ['да', 'нет']
            elif f.type == 'ref' and f.ref_directory:
                qs = ref_queryset(backend.org, f.ref_directory).filter(is_active=True)
                if f.ref_directory == 'categories':
                    values = [x.path_name() for x in qs if x.is_selectable() and x.is_leaf()]
                else:
                    values = [str(x) if f.ref_directory != 'models' else x.article for x in qs[:1000]]
                    if f.ref_directory == 'sizes':
                        values = [x.value if x.system == 'RU' else f'{x.value} {x.system}' for x in qs[:1000]]
        if values:
            list_col += 1
            letter = get_column_letter(list_col)
            lists.cell(row=1, column=list_col, value=c.label).font = Font(bold=True)
            for j, v in enumerate(values, start=2):
                lists.cell(row=j, column=list_col, value=v)
            dv = DataValidation(type='list', formula1=f"='Списки'!${letter}$2:${letter}${len(values) + 1}",
                                allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
            dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}10001')
    ws.freeze_panes = 'A2'
    lists.sheet_state = 'hidden' if list_col else 'visible'

    help_ws.append([f'Импорт: справочник «{backend.directory.name}»'])
    help_ws['A1'].font = Font(bold=True, size=13)
    help_ws.append([])
    for line in [
        'Заполните лист «Данные»: одна строка — одна запись. Колонки со звёздочкой * обязательны.',
        'Принимаются XLSX и CSV (UTF-8, разделитель «;» или «,»), до 10 000 строк за раз.',
        'Значения ссылок ищутся по коду или названию без учёта регистра.',
        'Несколько значений в одной ячейке разделяйте точкой с запятой: «королевское; кружевное».',
        'Даты — ДД.ММ.ГГГГ. Да/нет — «да» или «нет». Деньги — числом, без пробелов и знака ₸.',
        'При загрузке сначала показывается предпросмотр: сколько записей будет создано, обновлено,'
        ' пропущено, и список ошибок с номером строки.',
    ]:
        help_ws.append([line])
    if backend.code == 'items':
        help_ws.append(['Экземпляры: модель находится по колонке «Модель: артикул» или создаётся из колонок '
                        '«Модель: …». Так переносится весь каталог одной загрузкой.'])
    help_ws.append([])
    help_ws.append(['Колонка', 'Обязательная', 'Тип', 'Подсказка'])
    for c in help_ws[help_ws.max_row]:
        c.font = Font(bold=True)
    tnames = dict(m.FIELD_TYPE_CHOICES)
    for c in cols:
        d = c.as_dict()
        t = tnames.get(d['type'], d['type']) + (' (несколько через «;»)' if d['multiple'] else '')
        help_ws.append([c.label, 'да' if c.required else '', t, d['hint']])
    help_ws.column_dimensions['A'].width = 34
    help_ws.column_dimensions['C'].width = 30
    help_ws.column_dimensions['D'].width = 70
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
