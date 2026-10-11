"""
REST API модуля «Мастер-данные» (§8: все операции доступны через API).

/api/md/d/<справочник>/…  — единый набор эндпоинтов для всех справочников
(системных и пользовательских): список с фильтрами, карточка, создание,
изменение, архив/восстановление, история, массовые действия, экспорт,
шаблон импорта, значения для выпадающих списков.
"""
import csv
import io
import re
import uuid

from django.conf import settings
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.db import transaction
from django.db.models import Q
from django.http import FileResponse, Http404, HttpResponse
from django.utils import timezone
from rest_framework import exceptions, status
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.models import Membership, Organization

from .. import audit, fieldtypes, importer, labels, photos
from .. import models as m
from ..backends import get_backend
from ..registry import REGISTRY
from ..seed import ensure_role_policies
from .context import md
from .serializers import duplicate_warnings, photo_urls

PAGE_MAX = 200


def ok(data=None, status_code=200, **extra):
    payload = {'ok': True}
    if data is not None:
        payload['data'] = data
    payload.update(extra)
    return Response(payload, status=status_code)


class MDView(APIView):
    @property
    def ctx(self):
        return md(self.request)

    def backend(self, code, action='view'):
        ctx = self.ctx
        b = get_backend(ctx.org, code, ctx.user, ctx.policy(code))
        if not b.directory.is_active and action != 'view':
            raise exceptions.PermissionDenied('Справочник в архиве')
        ctx.require(code, action)
        return b

    def body(self):
        return self.request.data if isinstance(self.request.data, dict) else {}


# ============================================================
# Список / создание
# ============================================================
class DictListView(MDView):
    """GET — список (поиск q, фильтры по полям, ordering, page, page_size,
    archived=1|only); POST — создание записи."""

    def get(self, request, code):
        b = self.backend(code)
        params = request.query_params
        qs = b.apply_ordering(b.apply_filters(b.base_qs(), params), params.get('ordering'))
        try:
            page = max(1, int(params.get('page') or 1))
            size = min(PAGE_MAX, max(1, int(params.get('page_size') or 50)))
        except ValueError:
            page, size = 1, 50
        total = qs.count()
        objs = list(qs[(page - 1) * size: page * size])
        return Response({'ok': True, 'data': b.represent_many(objs, 'list'), 'total': total,
                         'page': page, 'page_size': size})

    def post(self, request, code):
        b = self.backend(code, 'create')
        obj, warnings = b.save(self.body(), source=_source(request))
        return ok(b.represent(obj), status.HTTP_201_CREATED, warnings=warnings)


def _source(request):
    return 'ui' if request.META.get('HTTP_X_CLIENT') == 'md-ui' else 'api'


class DictDetailView(MDView):
    def get(self, request, code, pk):
        b = self.backend(code)
        return ok(b.represent(b.get(pk)))

    def put(self, request, code, pk):
        return self._update(request, code, pk, partial=False)

    def patch(self, request, code, pk):
        return self._update(request, code, pk, partial=True)

    def _update(self, request, code, pk, partial):
        b = self.backend(code, 'update')
        obj = b.get(pk)
        obj, warnings = b.save(self.body(), obj, partial=partial, source=_source(request))
        return ok(b.represent(obj), warnings=warnings)

    def delete(self, request, code, pk):
        b = self.backend(code, 'archive')
        result, refs = b.delete(b.get(pk))
        msg = 'Запись удалена' if result == 'deleted' else 'На запись есть ссылки — она перенесена в архив'
        return ok({'result': result, 'references': refs}, message=msg)


class DictArchiveView(MDView):
    def post(self, request, code, pk):
        b = self.backend(code, 'archive')
        return ok(b.represent(b.archive(b.get(pk), _source(request))))


class DictRestoreView(MDView):
    def post(self, request, code, pk):
        b = self.backend(code, 'archive')
        return ok(b.represent(b.restore(b.get(pk), _source(request))))


class DictHistoryView(MDView):
    """МД-22: дата, пользователь, поле, старое и новое значение."""

    def get(self, request, code, pk):
        b = self.backend(code)
        obj = b.get(pk)
        p = self.ctx.policy(code)
        if p.history == 'none' or (p.history == 'own' and obj.created_by_id != self.ctx.user.pk):
            raise exceptions.PermissionDenied('Нет прав на просмотр истории этой записи')
        hidden = p.hidden
        rows = m.AuditLog.objects.filter(organization=self.ctx.org, entity=code, record_id=str(obj.pk)) \
            .select_related('user')[:500]
        actions = dict(m.AuditLog.ACTION_CHOICES)
        data = [{'at': r.at, 'user': r.user.name if r.user_id else '—', 'action': r.action,
                 'action__display': actions.get(r.action, r.action), 'source': r.source,
                 'changes': [c for c in r.changes if c.get('field') not in hidden]} for r in rows]
        return ok(data)


# ============================================================
# Метаданные и выпадающие списки
# ============================================================
SPECIAL_FILTERS = {
    'models': [
        {'code': 'size', 'name': 'Размер (есть экземпляр)', 'type': 'ref', 'ref_directory': 'sizes', 'multiple': True},
        {'code': 'salon', 'name': 'Салон', 'type': 'ref', 'ref_directory': 'salons', 'multiple': True},
        {'code': 'stage', 'name': 'Стадия', 'type': 'ref', 'ref_directory': 'stages', 'multiple': True},
        {'code': 'status', 'name': 'Статус экземпляра', 'type': 'choice', 'multiple': True,
         'options': [{'value': v, 'label': lbl} for v, lbl in m.ITEM_STATUS_CHOICES]},
        {'code': 'has_photo', 'name': 'Есть фото', 'type': 'bool'},
    ],
    'items': [
        {'code': 'category', 'name': 'Категория', 'type': 'ref', 'ref_directory': 'categories', 'multiple': True},
        {'code': 'events', 'name': 'Мероприятия', 'type': 'ref', 'ref_directory': 'events', 'multiple': True},
        {'code': 'styles', 'name': 'Виды', 'type': 'ref', 'ref_directory': 'styles', 'multiple': True},
        {'code': 'operations', 'name': 'Операции', 'type': 'choice', 'multiple': True,
         'options': [{'value': v, 'label': lbl} for v, lbl in m.OPERATION_CHOICES]},
        {'code': 'rent_price', 'name': 'Цена аренды, ₸', 'type': 'money'},
        {'code': 'available', 'name': 'Свободен сейчас', 'type': 'bool'},
        {'code': 'has_photo', 'name': 'Есть фото', 'type': 'bool'},
    ],
}


def field_meta(f, policy, backend):
    readonly = f.readonly or (policy is not None and f.code in policy.readonly)
    return {'code': f.code, 'name': f.name, 'type': f.type, 'required': f.required, 'unique': f.unique,
            'default': f.default, 'hint': f.hint, 'order': f.order, 'group': f.group or 'Основное',
            'show_in_list': f.show_in_list, 'searchable': f.searchable, 'filterable': f.filterable,
            'options': f.options, 'multiple': f.multiple, 'ref_directory': f.ref_directory,
            'validation': f.validation, 'readonly': readonly, 'is_system': f.is_system, 'id': str(f.id)}


class DictMetaView(MDView):
    def get(self, request, code):
        b = self.backend(code)
        d = b.directory
        p = self.ctx.policy(code)
        spec = REGISTRY.get(code)
        fields = [field_meta(f, p, b) for f in b.visible_fields()]
        filters = [f for f in fields if f['filterable']]
        filters += [x for x in SPECIAL_FILTERS.get(code, []) if x['code'] not in p.hidden]
        if b.supports_photo and code not in SPECIAL_FILTERS:
            filters.append({'code': 'has_photo', 'name': 'Есть фото', 'type': 'bool'})
        pref = m.UserPreference.objects.filter(organization=self.ctx.org, user=self.ctx.user,
                                               key=f'columns:{code}').first()
        return ok({
            'code': code, 'name': d.name, 'description': d.description, 'group': d.menu_group, 'icon': d.icon,
            'is_system': d.is_system, 'hierarchical': b.hierarchical, 'supports_photo': b.supports_photo,
            'supports_archive': d.supports_archive, 'display_field': d.display_field,
            'tabs': spec.tabs if spec else (['photos', 'history'] if b.supports_photo else ['history']),
            'fields': fields, 'filters': filters, 'permissions': p.as_dict(),
            'import_keys': [{'value': k, 'label': importer.KEY_LABELS.get(k, k)} for k in importer.key_fields(b)],
            'columns_pref': pref.value if pref else None,
            'tile_view': code == 'models',
        })


class DictOptionsView(MDView):
    """Значения для выпадающих списков: только активные и доступные к выбору
    (архивные и выключенные разделы скрыты, МД-6, МД-21)."""

    def get(self, request, code):
        b = self.backend(code)
        return ok(b.options(request.query_params))


# ============================================================
# Экспорт (МД-24), массовые действия (МД-23), шаблон импорта (МД-25)
# ============================================================
def export_table(b, objs):
    fields = [f for f in b.visible_fields() if f.type not in ('photo', 'file')
              and not (f.type == 'computed' and f.code not in ('article', 'kit'))]
    header = [f.name for f in fields]
    rows = []
    for o in objs:
        snap = audit.snapshot(o, b.all_fields())
        row = []
        for f in fields:
            if f.code == 'article' and isinstance(o, m.Item):
                row.append(o.article)
            elif f.code == 'kit' and isinstance(o, m.Item):
                link = o.active_kit_link
                row.append(link.kit.code if link else '')
            else:
                row.append(snap.get(f.code, ''))
        if not o.is_active:
            row.append('архив')
        rows.append(row)
    if any(not o.is_active for o in objs):
        header.append('Статус записи')
    return header, rows


def file_response(content, filename, content_type):
    resp = HttpResponse(content, content_type=content_type)
    from urllib.parse import quote
    resp['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"
    return resp


def render_export(b, objs, fmt):
    header, rows = export_table(b, objs)
    stamp = timezone.localtime().strftime('%Y%m%d_%H%M')
    name = f'{b.directory.name}_{stamp}'
    if fmt == 'csv':
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=';')
        w.writerow(header)
        w.writerows(rows)
        return file_response('﻿' + buf.getvalue(), f'{name}.csv', 'text/csv; charset=utf-8')
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = b.directory.name[:30]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = 'A2'
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(45, len(h) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    return file_response(buf.getvalue(), f'{name}.xlsx',
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class DictExportView(MDView):
    def get(self, request, code):
        b = self.backend(code, 'export')
        params = request.query_params
        qs = b.apply_ordering(b.apply_filters(b.base_qs(), params), params.get('ordering'))
        return render_export(b, list(qs[:50000]), params.get('format', 'xlsx'))


class DictTemplateView(MDView):
    def get(self, request, code):
        b = self.backend(code, 'import')
        return file_response(importer.build_template(b), f'Шаблон_{b.directory.name}.xlsx',
                             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class DictBulkView(MDView):
    """МД-23: изменить значение поля, архивировать, восстановить,
    экспортировать, напечатать бирки — для выбранных записей."""

    def post(self, request, code):
        body = self.body()
        action = body.get('action')
        ids = [str(x) for x in (body.get('ids') or [])]
        if not ids:
            raise exceptions.ValidationError({'ids': ['Выберите записи']})
        need = {'update': 'update', 'archive': 'archive', 'restore': 'archive', 'export': 'export',
                'labels': 'view'}.get(action)
        if need is None:
            raise exceptions.ValidationError({'action': ['Неизвестное действие']})
        b = self.backend(code, need)
        objs = list(b.apply_filters(b.base_qs(), {'ids': ','.join(ids), 'archived': '1'}))
        if action == 'export':
            return render_export(b, objs, body.get('format', 'xlsx'))
        if action == 'labels':
            if code != 'items':
                raise exceptions.ValidationError({'action': ['Бирки печатаются для экземпляров']})
            return file_response(labels.render_labels(objs, settings.APP_BASE_URL), 'Ярлыки.pdf', 'application/pdf')
        done, errors = 0, []
        field = body.get('field')
        if action == 'update':
            f = b.field(field)
            if f is None or f.readonly or f.type == 'computed':
                raise exceptions.ValidationError({'field': ['Поле нельзя изменить массово']})
        for o in objs:
            try:
                with transaction.atomic():
                    if action == 'update':
                        b.save({field: body.get('value')}, o, partial=True, source=_source(request))
                    elif action == 'archive':
                        b.archive(o)
                    elif action == 'restore':
                        b.restore(o)
                done += 1
            except exceptions.ValidationError as e:
                detail = e.detail
                msg = '; '.join(str(x) for v in (detail.values() if isinstance(detail, dict) else [detail])
                                for x in (v if isinstance(v, list) else [v]))
                errors.append({'id': str(o.pk), 'title': str(o), 'message': msg})
        return ok({'done': done, 'errors': errors})


# ============================================================
# Товар: экземпляры модели (МД-16), бирки (МД-17), фото (МД-19)
# ============================================================
class ModelAddItemsView(MDView):
    """Создаёт сразу несколько экземпляров модели: размеры × количество,
    общие поля (салон, поставщик, дата, стоимость) заполняются один раз."""

    def post(self, request, pk):
        mb = self.backend('models')
        model = mb.get(pk)
        ib = self.backend('items', 'create')
        body = self.body()
        lines = body.get('sizes') or []
        if not lines:
            raise exceptions.ValidationError({'sizes': ['Укажите размеры и количество']})
        common = {k: body.get(k) for k in ('salon', 'supplier', 'received_at', 'cost', 'condition', 'stage',
                                           'status', 'note') if body.get(k) not in (None, '')}
        total = 0
        for ln in lines:
            try:
                qty = int(ln.get('qty') or 0)
            except (TypeError, ValueError):
                raise exceptions.ValidationError({'sizes': ['Количество — целое число']})
            if qty < 0 or qty > 100:
                raise exceptions.ValidationError({'sizes': ['Количество — от 1 до 100 по каждому размеру']})
            total += qty
        if total == 0:
            raise exceptions.ValidationError({'sizes': ['Количество должно быть больше 0']})
        created = []
        with transaction.atomic():
            for ln in lines:
                for _ in range(int(ln.get('qty') or 0)):
                    data = dict(common, model=str(model.pk))
                    if ln.get('size'):
                        data['size'] = ln['size']
                    obj, _w = ib.save(data, source=_source(request))
                    created.append(obj)
        return ok(ib.represent_many(created, 'list'), status.HTTP_201_CREATED)


class LabelsView(MDView):
    """PDF с ярлыками 58×40 мм: ?ids=… (через запятую)."""

    def get(self, request):
        b = self.backend('items')
        ids = request.query_params.get('ids', '')
        objs = list(b.apply_filters(b.base_qs(), {'ids': ids, 'archived': '1'}))
        if not objs:
            raise exceptions.NotFound('Экземпляры не найдены')
        resp = HttpResponse(labels.render_labels(objs, settings.APP_BASE_URL), content_type='application/pdf')
        resp['Content-Disposition'] = 'inline; filename="labels.pdf"'
        return resp


def _photo_owner(b, obj):
    if isinstance(obj, m.ProductModel):
        return {'product_model': obj}
    if isinstance(obj, m.Item):
        return {'item': obj}
    return {'record': obj}


class PhotosView(MDView):
    def get(self, request, code, pk):
        b = self.backend(code)
        if not b.supports_photo:
            raise exceptions.NotFound('Справочник без фото')
        obj = b.get(pk)
        qs = m.Photo.objects.filter(organization=self.ctx.org, **_photo_owner(b, obj)).select_related('photo_type')
        return ok([photo_urls(p) for p in qs.order_by('-is_main', 'order', 'created_at')])

    def post(self, request, code, pk):
        b = self.backend(code, 'update')
        if not b.supports_photo:
            raise exceptions.NotFound('Справочник без фото')
        obj = b.get(pk)
        files = request.FILES.getlist('files') or request.FILES.getlist('file')
        if not files:
            raise exceptions.ValidationError({'file': ['Выберите файлы']})
        ptype = None
        if request.data.get('photo_type'):
            ptype = m.PhotoType.objects.filter(organization=self.ctx.org, pk=request.data['photo_type']).first()
        out = []
        with transaction.atomic():
            for f in files:
                p = photos.add_photo(self.ctx.org, self.ctx.user, f.read(), f.name, photo_type=ptype,
                                     **_photo_owner(b, obj))
                out.append(photo_urls(p))
            audit.log(obj, code, 'update', self.ctx.user,
                      [{'field': 'photos', 'label': 'Фото', 'old': '', 'new': f'+{len(files)} шт.'}], _source(request))
        return ok(out, status.HTTP_201_CREATED)


class PhotoDetailView(MDView):
    def _get(self, pk):
        p = m.Photo.objects.filter(organization=self.ctx.org, pk=pk).first()
        if p is None:
            raise exceptions.NotFound('Фото не найдено')
        code = 'models' if p.product_model_id else 'items' if p.item_id else p.record.directory.code
        self.ctx.require(code, 'update')
        return p

    def patch(self, request, pk):
        p = self._get(pk)
        body = self.body()
        owner = dict(product_model_id=p.product_model_id, item_id=p.item_id, record_id=p.record_id)
        with transaction.atomic():
            if 'photo_type' in body:
                p.photo_type = m.PhotoType.objects.filter(organization=self.ctx.org, pk=body['photo_type']).first() \
                    if body['photo_type'] else None
            if 'order' in body:
                p.order = int(body['order'])
            if body.get('is_main'):
                m.Photo.objects.filter(**owner).update(is_main=False)
                p.is_main = True
            p.save()
        return ok(photo_urls(p))

    def delete(self, request, pk):
        photos.delete_photo(self._get(pk))
        return ok()


class PhotosReorderView(MDView):
    def post(self, request):
        ids = self.body().get('ids') or []
        qs = {str(p.pk): p for p in m.Photo.objects.filter(organization=self.ctx.org, pk__in=ids)}
        for i, pid in enumerate(ids):
            p = qs.get(str(pid))
            if p:
                code = 'models' if p.product_model_id else 'items' if p.item_id else p.record.directory.code
                self.ctx.require(code, 'update')
                p.order = i + 1
                p.save(update_fields=['order'])
        return ok()


class PhotoFileView(APIView):
    """Отдаёт файл фото по подписанной ссылке с ограниченным сроком (§8)."""
    authentication_classes = []
    permission_classes = [AllowAny]

    def get(self, request, pk, variant):
        if variant not in ('file', 'thumb_300', 'thumb_1200'):
            raise Http404
        if not photos.check_signature(pk, variant, request.GET.get('exp'), request.GET.get('sig')):
            return HttpResponse('Ссылка недействительна или истекла', status=403)
        p = m.Photo.objects.filter(pk=pk).first()
        f = getattr(p, variant) if p else None
        if not f:
            raise Http404
        ext = (f.name.rsplit('.', 1)[-1] or '').lower()
        ctype = {'jpg': 'image/jpeg', 'jpeg': 'image/jpeg', 'png': 'image/png', 'webp': 'image/webp',
                 'heic': 'image/heic', 'heif': 'image/heif'}.get(ext, 'application/octet-stream')
        resp = FileResponse(f.open('rb'), content_type=ctype)
        resp['Cache-Control'] = 'private, max-age=600'
        return resp


class FileUploadView(MDView):
    """Файлы для полей типа «Файл» пользовательских справочников."""

    def post(self, request):
        f = request.FILES.get('file')
        if not f:
            raise exceptions.ValidationError({'file': ['Выберите файл']})
        if f.size > 20 * 1024 * 1024:
            raise exceptions.ValidationError({'file': ['Файл больше 20 МБ']})
        path = default_storage.save(f'md/{self.ctx.org.pk}/files/{uuid.uuid4().hex}_{f.name[-80:]}', f)
        return ok({'name': f.name, 'path': path})

    def get(self, request):
        path = request.query_params.get('path', '')
        if not path.startswith(f'md/{self.ctx.org.pk}/files/') or '..' in path or not default_storage.exists(path):
            raise exceptions.NotFound('Файл не найден')
        return FileResponse(default_storage.open(path, 'rb'), as_attachment=True,
                            filename=path.rsplit('/', 1)[-1].split('_', 1)[-1])


# ============================================================
# Клиенты: дубли (МД-20); поиск в шапке (§7)
# ============================================================
class ClientDuplicatesView(MDView):
    def get(self, request):
        self.backend('clients')
        q = request.query_params
        phone = q.get('phone') or ''
        try:
            from ..validators import normalize_phone
            phone = normalize_phone(phone) if phone else ''
        except Exception:
            pass
        return ok(duplicate_warnings(self.ctx.org, phone, q.get('full_name'), q.get('exclude') or None))


class SearchView(MDView):
    """Поиск по артикулу, бирке, наименованию, телефону и ИИН. Если найдена
    одна запись — UI открывает её сразу."""

    def get(self, request):
        q = (request.query_params.get('q') or '').strip()
        if len(q) < 2:
            return ok({'results': [], 'single': None})
        results = []
        ctx = self.ctx
        for code in ('items', 'models', 'kits', 'clients', 'counterparties'):
            if not ctx.policy(code).can['view']:
                continue
            b = get_backend(ctx.org, code, ctx.user, ctx.policy(code))
            qs = b.base_qs().filter(is_active=True)
            exact = Q()
            if code == 'items':
                exact = Q(tag__iexact=q) | Q(barcode=q)
                mt = re.match(r'^(.+)-([^-]+)$', q)
                if mt:
                    exact |= Q(model__article__iexact=mt.group(1), size__value__iexact=mt.group(2))
            elif code == 'models':
                exact = Q(article__iexact=q)
            elif code == 'kits':
                exact = Q(code__iexact=q)
            elif code == 'clients':
                d = re.sub(r'\D', '', q)
                if re.fullmatch(r'\d{12}', d):
                    from ..crypto import blind_hash
                    exact = Q(iin_hash=blind_hash(d))
                elif len(d) >= 10:
                    exact = Q(phone_digits__endswith=d[-10:])
            elif code == 'counterparties':
                exact = Q(bin_iin=q)
            found = list(qs.filter(exact)[:5]) if exact else []
            if not found:
                found = list(qs.filter(b.search_q(q))[:5])
            for o in found:
                sub = ''
                if code == 'items':
                    sub = f'{o.article} · {o.model.name} · {o.stage.name} · {o.salon.name}'
                elif code == 'models':
                    sub = o.category.path_name()
                elif code == 'clients':
                    from ..validators import format_phone
                    sub = format_phone(o.phone)
                elif code == 'counterparties':
                    sub = o.get_kind_display()
                title = o.tag if code == 'items' else str(o)
                results.append({'dict': code, 'dict_name': b.directory.name, 'id': str(o.pk), 'title': title,
                                'subtitle': sub, 'exact': bool(exact) and o in list(qs.filter(exact)[:5])})
        exact_hits = [r for r in results if r['exact']]
        single = exact_hits[0] if len(exact_hits) == 1 else (results[0] if len(results) == 1 else None)
        return ok({'results': results, 'single': single})


# ============================================================
# Конструктор справочников (МД-1…12)
# ============================================================
CODE_RE = re.compile(r'^[a-z][a-z0-9_]{1,40}$')
RESERVED_FIELD_CODES = {'id', 'code', 'parent', 'is_active', 'extra', 'data', 'organization', 'created_at',
                        'created_by', 'updated_at', 'updated_by', 'photos', 'q', 'page', 'page_size', 'ordering',
                        'archived', 'ids', 'has_photo', 'format', 'org'}
DIR_EDITABLE_SYSTEM = {'name', 'description', 'icon', 'menu_enabled', 'order'}
DIR_EDITABLE_CUSTOM = DIR_EDITABLE_SYSTEM | {'menu_group', 'is_hierarchical', 'display_field', 'supports_photo',
                                             'supports_archive'}
FIELD_EDITABLE_SYSTEM = {'name', 'order', 'group', 'hint', 'show_in_list', 'searchable', 'filterable'}
FIELD_EDITABLE = FIELD_EDITABLE_SYSTEM | {'type', 'required', 'unique', 'default', 'options', 'multiple',
                                          'ref_directory', 'validation'}


def directory_dict(d, ctx=None):
    out = {'id': str(d.id), 'code': d.code, 'name': d.name, 'description': d.description,
           'menu_group': d.menu_group, 'menu_group__display': d.get_menu_group_display(), 'icon': d.icon,
           'is_system': d.is_system, 'is_hierarchical': d.is_hierarchical, 'display_field': d.display_field,
           'supports_photo': d.supports_photo, 'supports_archive': d.supports_archive,
           'menu_enabled': d.menu_enabled, 'order': d.order, 'is_active': d.is_active}
    if ctx is not None:
        out['permissions'] = ctx.policy(d.code).as_dict()
    return out


class DirectoriesView(MDView):
    def get(self, request):
        ctx = self.ctx
        qs = m.Directory.objects.filter(organization=ctx.org)
        if request.query_params.get('archived') not in ('1', 'true'):
            qs = qs.filter(is_active=True)
        data = [directory_dict(d, ctx) for d in qs if ctx.policy(d.code).can['view'] or ctx.is_manager]
        return ok(data)

    def post(self, request):
        ctx = self.ctx
        ctx.require_manager()
        body = self.body()
        code = (body.get('code') or '').strip().lower()
        name = (body.get('name') or '').strip()
        errors = {}
        if not CODE_RE.match(code):
            errors['code'] = ['Код — латиница, цифры и «_», начинается с буквы (2–40 символов)']
        elif code in REGISTRY or m.Directory.objects.filter(organization=ctx.org, code=code).exists():
            errors['code'] = ['Справочник с таким кодом уже есть']
        if not name:
            errors['name'] = ['Обязательное поле']
        if errors:
            raise exceptions.ValidationError(errors)
        with transaction.atomic():
            d = m.Directory.objects.create(
                organization=ctx.org, code=code, name=name, description=body.get('description', ''),
                menu_group=body.get('menu_group') or 'custom', icon=body.get('icon') or '📁',
                is_hierarchical=bool(body.get('is_hierarchical')), supports_photo=bool(body.get('supports_photo')),
                supports_archive=body.get('supports_archive', True) is not False,
                display_field=body.get('display_field') or 'name', created_by=ctx.user)
            if d.display_field == 'name':
                m.DirectoryField.objects.create(organization=ctx.org, directory=d, code='name', name='Наименование',
                                                type='string', required=True, searchable=True, order=10)
            ensure_role_policies(ctx.org, [code])
            audit.structure_log(ctx.org, code, 'directory_create', ctx.user, name=name)
        return ok(directory_dict(d, ctx), status.HTTP_201_CREATED)


def _directory(ctx, code):
    d = m.Directory.objects.filter(organization=ctx.org, code=code).first()
    if d is None:
        raise exceptions.NotFound('Справочник не найден')
    return d


class DirectoryDetailView(MDView):
    def get(self, request, code):
        d = _directory(self.ctx, code)
        return ok(directory_dict(d, self.ctx))

    def patch(self, request, code):
        ctx = self.ctx
        ctx.require_manager()
        d = _directory(ctx, code)
        allowed = DIR_EDITABLE_SYSTEM if d.is_system else DIR_EDITABLE_CUSTOM
        body = self.body()
        changes = {}
        for k, v in body.items():
            if k not in allowed:
                if k in DIR_EDITABLE_CUSTOM:
                    raise exceptions.ValidationError({k: ['У системного справочника это свойство не меняется']})
                continue
            old = getattr(d, k)
            if old != v:
                changes[k] = {'old': old, 'new': v}
                setattr(d, k, v)
        if 'display_field' in changes and not d.fields.filter(code=d.display_field).exists():
            raise exceptions.ValidationError({'display_field': ['Нет такого поля']})
        if not d.name.strip():
            raise exceptions.ValidationError({'name': ['Обязательное поле']})
        d.updated_by = ctx.user
        d.save()
        if changes:
            audit.structure_log(ctx.org, code, 'directory_update', ctx.user, changes=changes)
        return ok(directory_dict(d, ctx))

    def delete(self, request, code):
        """МД-5: удалить можно только пустой справочник без ссылок на него;
        иначе — только архивировать. Системные удалить нельзя (МД-4)."""
        ctx = self.ctx
        ctx.require_manager()
        d = _directory(ctx, code)
        if d.is_system:
            raise exceptions.ValidationError('Системный справочник удалить нельзя — его можно отключить из меню')
        has_records = d.records.exists()
        refs = m.DirectoryField.objects.filter(organization=ctx.org, type='ref', ref_directory=code,
                                               is_archived=False).exclude(directory=d).select_related('directory')
        if has_records or refs.exists() or request.query_params.get('archive') in ('1', 'true'):
            d.is_active = False
            d.menu_enabled = False
            d.save()
            audit.structure_log(ctx.org, code, 'directory_archive', ctx.user)
            why = []
            if has_records:
                why.append('в справочнике есть записи')
            if refs.exists():
                why.append('на него ссылаются поля: ' + ', '.join(f'{f.directory.name}.{f.name}' for f in refs[:5]))
            return ok({'result': 'archived'}, message='Справочник перенесён в архив' + (
                f' ({"; ".join(why)})' if why else ''))
        audit.structure_log(ctx.org, code, 'directory_delete', ctx.user, name=d.name)
        m.RolePolicy.objects.filter(organization=ctx.org, directory_code=code).delete()
        d.delete()
        return ok({'result': 'deleted'}, message='Справочник удалён')


class DirectoryRestoreView(MDView):
    def post(self, request, code):
        ctx = self.ctx
        ctx.require_manager()
        d = _directory(ctx, code)
        d.is_active, d.menu_enabled = True, True
        d.save()
        audit.structure_log(ctx.org, code, 'directory_restore', ctx.user)
        return ok(directory_dict(d, ctx))


class DirectoryCopyView(MDView):
    """МД-3: копия справочника со структурой, но без данных."""

    def post(self, request, code):
        ctx = self.ctx
        ctx.require_manager()
        src = _directory(ctx, code)
        body = self.body()
        new_code = (body.get('code') or '').strip().lower()
        if not CODE_RE.match(new_code):
            raise exceptions.ValidationError({'code': ['Код — латиница, цифры и «_», начинается с буквы']})
        if new_code in REGISTRY or m.Directory.objects.filter(organization=ctx.org, code=new_code).exists():
            raise exceptions.ValidationError({'code': ['Справочник с таким кодом уже есть']})
        with transaction.atomic():
            d = m.Directory.objects.create(
                organization=ctx.org, code=new_code, name=body.get('name') or f'{src.name} (копия)',
                description=src.description, menu_group='custom' if src.is_system else src.menu_group, icon=src.icon,
                is_hierarchical=src.is_hierarchical, supports_photo=src.supports_photo,
                supports_archive=src.supports_archive, created_by=ctx.user,
                display_field=src.display_field if not src.is_system else 'name')
            fields = list(src.fields.filter(is_archived=False))
            for f in fields:
                if f.type == 'computed':
                    continue
                ftype = f.type
                multiple = f.multiple
                m.DirectoryField.objects.create(
                    organization=ctx.org, directory=d, code=f.code, name=f.name, type=ftype, required=f.required,
                    unique=f.unique, default=f.default, hint=f.hint, order=f.order, group=f.group,
                    show_in_list=f.show_in_list, searchable=f.searchable, filterable=f.filterable,
                    options=f.options, multiple=multiple, ref_directory=f.ref_directory, validation=f.validation)
            if not d.fields.filter(code=d.display_field).exists():
                first = d.fields.order_by('order').first()
                d.display_field = first.code if first else 'name'
                d.save()
            ensure_role_policies(ctx.org, [new_code])
            audit.structure_log(ctx.org, new_code, 'directory_copy', ctx.user, source=code)
        return ok(directory_dict(d, ctx), status.HTTP_201_CREATED)


def _field_has_data(ctx, d, code):
    if d.is_system:
        spec = REGISTRY[d.code]
        qs = spec.model.objects.filter(organization=ctx.org, extra__has_key=code).exclude(**{f'extra__{code}': None})
        return qs.count()
    return m.DirectoryRecord.objects.filter(directory=d, data__has_key=code).exclude(**{f'data__{code}': None}).count()


def field_dict(f):
    return {'id': str(f.id), 'code': f.code, 'name': f.name, 'type': f.type, 'type__display': f.get_type_display(),
            'is_system': f.is_system, 'required': f.required, 'unique': f.unique, 'default': f.default,
            'hint': f.hint, 'order': f.order, 'group': f.group, 'show_in_list': f.show_in_list,
            'searchable': f.searchable, 'filterable': f.filterable, 'options': f.options, 'multiple': f.multiple,
            'ref_directory': f.ref_directory, 'validation': f.validation, 'readonly': f.readonly,
            'is_archived': f.is_archived}


def _validate_field_body(ctx, d, body, field=None):
    errors = {}
    types = {t for t, _ in m.FIELD_TYPE_CHOICES}
    if 'type' in body:
        if body['type'] not in types or body['type'] == 'computed':
            errors['type'] = ['Недопустимый тип поля (вычисляемые — только системные)']
    t = body.get('type', field.type if field else 'string')
    if 'name' in body and not str(body['name']).strip():
        errors['name'] = ['Обязательное поле']
    if t == 'choice':
        opts = body.get('options', field.options if field else [])
        norm = []
        for o in opts or []:
            if isinstance(o, str):
                o = {'value': o, 'label': o}
            if not str(o.get('label', '')).strip():
                continue
            norm.append({'value': str(o.get('value') or o['label']).strip(), 'label': str(o['label']).strip()})
        if not norm:
            errors['options'] = ['Укажите значения списка']
        body['options'] = norm
    if t == 'ref':
        rd = body.get('ref_directory', field.ref_directory if field else '')
        if not rd or not (rd in REGISTRY or m.Directory.objects.filter(organization=ctx.org, code=rd).exists()):
            errors['ref_directory'] = ['Выберите справочник для ссылки']
    v = body.get('validation')
    if v and v.get('mask'):
        try:
            re.compile(v['mask'])
        except re.error:
            errors['validation'] = ['Некорректная маска (регулярное выражение)']
    if body.get('multiple') and t not in ('choice', 'ref', 'string'):
        errors['multiple'] = ['Несколько значений — только для списков, ссылок и строк']
    if errors:
        raise exceptions.ValidationError(errors)


class FieldsView(MDView):
    def get(self, request, code):
        ctx = self.ctx
        d = _directory(ctx, code)
        qs = d.fields.all()
        if request.query_params.get('archived') not in ('1', 'true'):
            qs = qs.filter(is_archived=False)
        return ok([field_dict(f) for f in qs.order_by('order', 'created_at')])

    def post(self, request, code):
        ctx = self.ctx
        ctx.require_manager()
        d = _directory(ctx, code)
        body = dict(self.body())
        fcode = (body.get('code') or '').strip().lower()
        if not CODE_RE.match(fcode):
            raise exceptions.ValidationError({'code': ['Код — латиница, цифры и «_», начинается с буквы']})
        if fcode in RESERVED_FIELD_CODES or d.fields.filter(code=fcode).exists():
            raise exceptions.ValidationError({'code': ['Поле с таким кодом уже есть или код зарезервирован']})
        if d.is_system:
            model_fields = {f.name for f in REGISTRY[d.code].model._meta.get_fields()}
            if fcode in model_fields:
                raise exceptions.ValidationError({'code': ['Код совпадает с системным полем']})
        body.setdefault('type', 'string')
        _validate_field_body(ctx, d, body)
        last = d.fields.order_by('-order').values_list('order', flat=True).first() or 0
        vals = {k: body[k] for k in FIELD_EDITABLE if k in body}
        vals.setdefault('order', last + 10)
        vals.setdefault('name', fcode)
        f = m.DirectoryField.objects.create(organization=ctx.org, directory=d, code=fcode, is_system=False, **vals)
        audit.structure_log(ctx.org, code, 'field_add', ctx.user, field=fcode, name=f.name, type=f.type)
        return ok(field_dict(f), status.HTTP_201_CREATED)


class FieldDetailView(MDView):
    def _get(self, code, fid):
        d = _directory(self.ctx, code)
        f = d.fields.filter(pk=fid).first()
        if f is None:
            raise exceptions.NotFound('Поле не найдено')
        return d, f

    def patch(self, request, code, fid):
        ctx = self.ctx
        ctx.require_manager()
        d, f = self._get(code, fid)
        body = dict(self.body())
        allowed = FIELD_EDITABLE_SYSTEM if f.is_system else FIELD_EDITABLE
        bad = [k for k in body if k not in allowed and k in FIELD_EDITABLE]
        if bad:
            raise exceptions.ValidationError({bad[0]: ['Системное поле: можно менять только подпись, порядок, '
                                                       'группу и признаки списка/поиска/фильтра (МД-10)']})
        body = {k: v for k, v in body.items() if k in allowed}
        if 'type' in body and body['type'] != f.type:
            if _field_has_data(ctx, d, f.code) and (f.type, body['type']) not in fieldtypes.SAFE_TYPE_CHANGES:
                raise exceptions.ValidationError({'type': ['В поле есть данные: допускаются только безопасные '
                                                           'преобразования (целое → дробное, строка → текст)']})
        if not f.is_system:
            _validate_field_body(ctx, d, body, f)
        changes = {}
        for k, v in body.items():
            if getattr(f, k) != v:
                changes[k] = {'old': getattr(f, k), 'new': v}
                setattr(f, k, v)
        f.save()
        if changes:
            audit.structure_log(ctx.org, code, 'field_update', ctx.user, field=f.code, changes=changes)
        return ok(field_dict(f))

    def delete(self, request, code, fid):
        """МД-11: удаление поля с данными требует подтверждения и архивирует
        значения, а не стирает их."""
        ctx = self.ctx
        ctx.require_manager()
        d, f = self._get(code, fid)
        if f.is_system:
            raise exceptions.ValidationError('Системное поле нельзя удалить')
        cnt = _field_has_data(ctx, d, f.code)
        if cnt and request.query_params.get('confirm') not in ('1', 'true'):
            return Response({'ok': False, 'requires_confirm': True, 'records': cnt,
                             'error': f'В поле есть данные ({cnt} записей). Подтвердите удаление: значения '
                                      f'будут перенесены в архив, а не стёрты.'}, status=409)
        if cnt:
            f.is_archived = True
            f.save()
            audit.structure_log(ctx.org, code, 'field_archive', ctx.user, field=f.code, records=cnt)
            return ok({'result': 'archived'})
        audit.structure_log(ctx.org, code, 'field_delete', ctx.user, field=f.code)
        f.delete()
        return ok({'result': 'deleted'})


class FieldRestoreView(MDView):
    def post(self, request, code, fid):
        ctx = self.ctx
        ctx.require_manager()
        d = _directory(ctx, code)
        f = d.fields.filter(pk=fid).first()
        if f is None:
            raise exceptions.NotFound('Поле не найдено')
        f.is_archived = False
        f.save()
        audit.structure_log(ctx.org, code, 'field_restore', ctx.user, field=f.code)
        return ok(field_dict(f))


class FieldsReorderView(MDView):
    def post(self, request, code):
        ctx = self.ctx
        ctx.require_manager()
        d = _directory(ctx, code)
        ids = [str(x) for x in (self.body().get('ids') or [])]
        fields = {str(f.id): f for f in d.fields.all()}
        for i, fid in enumerate(ids):
            f = fields.get(fid)
            if f and f.order != (i + 1) * 10:
                f.order = (i + 1) * 10
                f.save(update_fields=['order'])
        audit.structure_log(ctx.org, code, 'fields_reorder', ctx.user, order=[fields[i].code for i in ids if i in fields])
        return ok()


class StructureLogView(MDView):
    def get(self, request, code=None):
        ctx = self.ctx
        ctx.require_manager()
        qs = m.StructureLog.objects.filter(organization=ctx.org).select_related('user')
        if code:
            qs = qs.filter(directory_code=code)
        return ok([{'at': r.at, 'user': r.user.name if r.user_id else '—', 'directory': r.directory_code,
                    'action': r.action, 'details': r.details} for r in qs[:300]])


# ============================================================
# Организация, роли, настройки пользователя
# ============================================================
def org_payload(ctx):
    org = ctx.org
    return {
        'organization': {'id': str(org.id), 'name': org.name, 'sections': org.sections()},
        'role': ctx.role, 'role__display': dict(Membership.ROLE_CHOICES)[ctx.role],
        'is_manager': ctx.is_manager,
        'user': {'id': ctx.user.pk, 'name': ctx.user.name, 'email': ctx.user.email},
        'memberships': [{'organization': str(x.organization_id), 'name': x.organization.name, 'role': x.role}
                        for x in Membership.objects.filter(user=ctx.user, is_active=True)
                        .select_related('organization')],
        'menu': [directory_dict(d, ctx) for d in m.Directory.objects.filter(organization=org, is_active=True)
                 if ctx.policy(d.code).can['view']],
    }


class OrgView(MDView):
    def get(self, request):
        return ok(org_payload(self.ctx))

    def patch(self, request):
        ctx = self.ctx
        ctx.require_manager()
        org = ctx.org
        body = self.body()
        settings_ = dict(org.settings or {})
        changes = {}
        if 'name' in body and str(body['name']).strip():
            changes['name'] = {'old': org.name, 'new': body['name']}
            org.name = str(body['name']).strip()
        if 'sections' in body:
            sec = dict(org.sections())
            for k, v in (body['sections'] or {}).items():
                if k in org.DEFAULT_SECTIONS:
                    if sec.get(k) != bool(v):
                        changes[f'section:{k}'] = {'old': sec.get(k), 'new': bool(v)}
                    sec[k] = bool(v)
            settings_['sections'] = sec
        org.settings = settings_
        org.save()
        if changes:
            audit.structure_log(org, '_organization', 'settings_update', ctx.user, changes=changes)
        return ok(org_payload(ctx))


class MembersView(MDView):
    """Участники организации и их роли (владелец/администратор)."""

    def get(self, request):
        ctx = self.ctx
        ctx.require_manager()
        qs = Membership.objects.filter(organization=ctx.org).select_related('user')
        return ok([{'id': x.pk, 'user': x.user.name, 'email': x.user.email, 'role': x.role, 'is_active': x.is_active}
                   for x in qs])

    def post(self, request):
        ctx = self.ctx
        ctx.require_manager()
        from django.contrib.auth import get_user_model
        body = self.body()
        role = body.get('role') or 'seller'
        if role not in dict(Membership.ROLE_CHOICES):
            raise exceptions.ValidationError({'role': ['Неизвестная роль']})
        if role == 'owner' and ctx.role != 'owner':
            raise exceptions.PermissionDenied('Назначить владельца может только владелец')
        user = get_user_model().objects.filter(email=(body.get('email') or '').strip().lower()).first()
        if user is None:
            raise exceptions.ValidationError({'email': ['Пользователь не зарегистрирован в системе']})
        mem, _ = Membership.objects.update_or_create(user=user, organization=ctx.org,
                                                     defaults={'role': role, 'is_active': True})
        audit.structure_log(ctx.org, '_members', 'member_set', ctx.user, email=user.email, role=role)
        return ok({'id': mem.pk, 'user': user.name, 'email': user.email, 'role': mem.role})


class OrganizationCreateView(MDView):
    """Регистрация новой организации: текущий пользователь становится
    владельцем; системные справочники создаются автоматически (МД-4)."""

    def post(self, request):
        name = (self.body().get('name') or '').strip()
        if len(name) < 2:
            raise exceptions.ValidationError({'name': ['Укажите название организации']})
        with transaction.atomic():
            org = Organization.objects.create(name=name)
            Membership.objects.create(user=request.user, organization=org, role=Membership.ROLE_OWNER)
        return ok({'id': str(org.id), 'name': org.name}, status.HTTP_201_CREATED)


POLICY_KEYS = ['can_view', 'can_create', 'can_update', 'can_archive', 'can_import', 'can_export']


class RolePoliciesView(MDView):
    """Матрица «роль × справочник × действие» + скрытые/read-only поля (§6)."""

    def get(self, request):
        ctx = self.ctx
        ctx.require_manager()
        codes = list(m.Directory.objects.filter(organization=ctx.org, is_active=True).values_list('code', flat=True))
        ensure_role_policies(ctx.org, codes)
        rows = m.RolePolicy.objects.filter(organization=ctx.org, directory_code__in=codes)
        return ok([{'role': r.role, 'directory_code': r.directory_code, **{k: getattr(r, k) for k in POLICY_KEYS},
                    'history': r.history, 'hidden_fields': r.hidden_fields, 'readonly_fields': r.readonly_fields}
                   for r in rows])

    def put(self, request):
        ctx = self.ctx
        ctx.require_manager()
        rows = self.body().get('rows') or []
        changed = 0
        for row in rows:
            role, code = row.get('role'), row.get('directory_code')
            if role == 'owner':
                continue  # владелец всегда имеет полные права
            r = m.RolePolicy.objects.filter(organization=ctx.org, role=role, directory_code=code).first()
            if r is None:
                continue
            for k in POLICY_KEYS:
                if k in row:
                    setattr(r, k, bool(row[k]))
            if row.get('history') in ('all', 'own', 'none'):
                r.history = row['history']
            for k in ('hidden_fields', 'readonly_fields'):
                if isinstance(row.get(k), list):
                    setattr(r, k, [str(x) for x in row[k]])
            r.save()
            changed += 1
        audit.structure_log(ctx.org, '_roles', 'policies_update', ctx.user, rows=changed)
        return ok({'changed': changed})


class PrefsView(MDView):
    """Настройки пользователя (колонки и их порядок — МД-13)."""

    def get(self, request, key):
        p = m.UserPreference.objects.filter(organization=self.ctx.org, user=self.ctx.user, key=key).first()
        return ok(p.value if p else None)

    def put(self, request, key):
        val = self.body().get('value')
        m.UserPreference.objects.update_or_create(organization=self.ctx.org, user=self.ctx.user, key=key[:100],
                                                  defaults={'value': val})
        return ok(val)


# ============================================================
# Импорт (МД-25…35)
# ============================================================
def job_dict(j, with_details=False):
    d = {'id': str(j.id), 'directory': j.directory_code, 'file': j.original_name, 'status': j.status,
         'status__display': j.get_status_display(), 'mode': j.mode, 'key_field': j.key_field,
         'total_rows': j.total_rows, 'processed_rows': j.processed_rows, 'counts': j.counts,
         'message': j.message, 'user': j.user.name if j.user_id else '', 'created_at': j.created_at,
         'finished_at': j.finished_at, 'has_error_file': bool(j.error_file), 'errors_count': len(j.errors or [])}
    if with_details:
        d.update({'errors': j.errors[:1000], 'mapping': j.mapping, 'headers': j.headers,
                  'skipped': (j.options or {}).get('skipped', [])[:200], 'options': {
                      k: v for k, v in (j.options or {}).items() if k != 'skipped'}})
    return d


class ImportUploadView(MDView):
    """Шаг 1: файл. Возвращает колонки файла, поля справочника,
    автоматическое сопоставление и сохранённые профили."""

    def post(self, request):
        code = request.data.get('directory')
        b = self.backend(code, 'import')
        f = request.FILES.get('file')
        if not f:
            raise exceptions.ValidationError({'file': ['Выберите файл']})
        raw = f.read()
        headers, rows = importer.read_table(raw, f.name)
        cols = importer.columns_for(b)
        mapping = importer.auto_map(headers, cols)
        job = m.ImportJob(organization=self.ctx.org, directory_code=code, user=self.ctx.user,
                          original_name=f.name[:255], headers=headers, total_rows=len(rows), mapping=mapping,
                          key_field=importer.key_fields(b)[0])
        job.file.save(f.name[-80:], ContentFile(raw), save=False)
        job.save()
        profiles = m.ImportProfile.objects.filter(organization=self.ctx.org, directory_code=code)
        sample = [[_jsonable(c) for c in r[1][:len(headers)]] for r in rows[:5]]
        return ok({'job': job_dict(job), 'headers': headers, 'sample': sample,
                   'columns': [c.as_dict() for c in cols], 'mapping': mapping,
                   'keys': [{'value': k, 'label': importer.KEY_LABELS.get(k, k)} for k in importer.key_fields(b)],
                   'profiles': [{'id': p.pk, 'name': p.name, 'mapping': p.mapping} for p in profiles]},
                  status.HTTP_201_CREATED)


def _jsonable(v):
    if v is None:
        return ''
    if hasattr(v, 'isoformat'):
        return v.isoformat()
    return v if isinstance(v, (int, float, str, bool)) else str(v)


def _job(ctx, pk):
    j = m.ImportJob.objects.filter(organization=ctx.org, pk=pk).first()
    if j is None:
        raise exceptions.NotFound('Загрузка не найдена')
    return j


class ImportJobView(MDView):
    def get(self, request, pk):
        return ok(job_dict(_job(self.ctx, pk), with_details=True))


class ImportPreviewView(MDView):
    """Шаг 2→3: сопоставление и режим → проверка всех строк (МД-30)."""

    def post(self, request, pk):
        ctx = self.ctx
        job = _job(ctx, pk)
        b = self.backend(job.directory_code, 'import')
        body = self.body()
        if job.status in ('running', 'previewing'):
            raise exceptions.ValidationError('Загрузка уже выполняется')
        mapping = {h: k for h, k in (body.get('mapping') or job.mapping or {}).items() if k}
        valid_keys = {c.key for c in importer.columns_for(b)}
        mapping = {h: k for h, k in mapping.items() if k in valid_keys and h in job.headers}
        if not mapping:
            raise exceptions.ValidationError({'mapping': ['Сопоставьте хотя бы одну колонку']})
        mode = body.get('mode') or 'create'
        if mode not in dict(m.ImportJob.MODE_CHOICES):
            raise exceptions.ValidationError({'mode': ['Неизвестный режим']})
        key = body.get('key_field') or importer.key_fields(b)[0]
        if key not in importer.key_fields(b):
            raise exceptions.ValidationError({'key_field': ['Недопустимый ключ сопоставления']})
        if mode in ('update', 'upsert') and key not in mapping.values():
            raise exceptions.ValidationError({'key_field': ['Колонка с ключом сопоставления не выбрана']})
        job.mapping, job.mode, job.key_field = mapping, mode, key
        job.options = {'create_missing_refs': bool(body.get('create_missing_refs'))}
        job.status, job.processed_rows = 'previewing', 0
        job.save()
        if body.get('profile_name'):
            m.ImportProfile.objects.update_or_create(
                organization=ctx.org, directory_code=job.directory_code, name=str(body['profile_name'])[:100],
                defaults={'mapping': mapping, 'created_by': ctx.user})
        from ..tasks import start_job
        how = start_job(job, apply=False, background=job.total_rows > settings.MD_IMPORT_BACKGROUND_ROWS)
        job.refresh_from_db()
        return ok(job_dict(job, with_details=True), background=how != 'sync')


class ImportRunView(MDView):
    """Шаг 4: запись. Целиком или «только корректные строки» (МД-31)."""

    def post(self, request, pk):
        ctx = self.ctx
        job = _job(ctx, pk)
        self.backend(job.directory_code, 'import')
        if job.status not in ('previewed', 'failed'):
            raise exceptions.ValidationError('Сначала выполните предпросмотр')
        opts = dict(job.options or {})
        opts['only_valid'] = bool(self.body().get('only_valid'))
        job.options = opts
        job.status, job.processed_rows = 'running', 0
        job.save()
        from ..tasks import start_job
        how = start_job(job, apply=True, background=job.total_rows > settings.MD_IMPORT_BACKGROUND_ROWS)
        job.refresh_from_db()
        return ok(job_dict(job, with_details=True), background=how != 'sync')


class ImportJobsView(MDView):
    """Журнал загрузок (МД-32): файл, пользователь, время, итоги."""

    def get(self, request):
        ctx = self.ctx
        qs = m.ImportJob.objects.filter(organization=ctx.org).select_related('user')
        if not ctx.is_manager:
            qs = qs.filter(user=ctx.user)
        if request.query_params.get('directory'):
            qs = qs.filter(directory_code=request.query_params['directory'])
        return ok([job_dict(j) for j in qs[:200]])


class ImportErrorFileView(MDView):
    def get(self, request, pk):
        job = _job(self.ctx, pk)
        if not job.error_file:
            raise exceptions.NotFound('Файла ошибок нет')
        return FileResponse(job.error_file.open('rb'), as_attachment=True,
                            filename=f'Ошибки_{job.original_name.rsplit(".", 1)[0]}.xlsx')


class ImportProfileDetailView(MDView):
    def delete(self, request, pk):
        p = m.ImportProfile.objects.filter(organization=self.ctx.org, pk=pk).first()
        if p is None:
            raise exceptions.NotFound('Профиль не найден')
        self.ctx.require(p.directory_code, 'import')
        p.delete()
        return ok()


class PhotoZipImportView(MDView):
    """МД-34: ZIP с фото, имя файла — артикул модели или бирка (СВ-117_1.jpg)."""

    def post(self, request):
        self.backend('models', 'import')
        f = request.FILES.get('file')
        if not f:
            raise exceptions.ValidationError({'file': ['Выберите ZIP-архив']})
        with transaction.atomic():
            report = photos.import_zip(self.ctx.org, self.ctx.user, f.read())
        return ok(report)
